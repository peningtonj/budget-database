"""
One-off export for manual data-quality review: every distinct (portfolio,
agency, outcome, program) combination as it literally appears in
program_expenses, alongside its canonicalized portfolio/agency, with each
year's "estimated actual" figure pivoted into its own column.

Grouped by the *raw* tuple (not the canonicalized one) deliberately: the
whole point is to let a human see the fragmentation itself -- e.g. "Foreign
Affairs", "DFAT" and "Foreign Affairs & Trade" showing up as three
adjacent rows (sorted by canonical portfolio) with disjoint year columns,
rather than hiding that behind an already-merged view.

Portfolio canonicalization is imported directly from the repo-root
portfolio_aliases module (shared with backend/measures/views.py and
build_db.py -- no more per-script copy to keep in sync).
"""
import os
import sqlite3
from collections import defaultdict, Counter

import openpyxl
from openpyxl.utils import get_column_letter

from portfolio_aliases import canon_portfolio

# Repo root (this file lives in pipeline/); programs.db sits one level up.
# The .xlsx is written there too, where .gitignore already excludes it.
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH = os.path.join(_ROOT, "programs.db")
OUT_PATH = os.path.join(_ROOT, "program_data_audit.xlsx")
YEARS = [f"{y}-{(y + 1) % 100:02d}" for y in range(2014, 2026)]  # 2014-15 .. 2025-26


# ---- agency canonicalization: formal_name per (edition, short_name) -------

con = sqlite3.connect(DB_PATH)
cur = con.cursor()
cur.execute("select edition, short_name, formal_name from agency_aliases")
formal_by_edition_agency = {(e, s): f for e, s, f in cur.fetchall()}


def main():
    cur.execute(
        """
        select edition, portfolio, agency, outcome_number, outcome_description,
               program_number, program_name, fiscal_year, amount_thousands
        from program_expenses
        where estimate_type = 'estimated_actual'
        """
    )
    rows = cur.fetchall()

    groups = {}
    for (edition, portfolio, agency, outcome_number, outcome_desc,
         program_number, program_name, fiscal_year, amount) in rows:
        key = (portfolio, agency, outcome_number, outcome_desc, program_number, program_name)
        g = groups.setdefault(key, {"years": {}, "formal_names": Counter(), "editions": set()})
        if fiscal_year in YEARS:
            g["years"].setdefault(fiscal_year, amount)  # keep first if a raw combo somehow repeats a year
        g["editions"].add(edition)
        fn = formal_by_edition_agency.get((edition, agency))
        if fn:
            g["formal_names"][fn] += 1

    print(f"distinct (portfolio, agency, outcome, program) rows: {len(groups)}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Program audit"

    headers = [
        "Portfolio", "Canonised Portfolio", "Agency", "Canonised Agency",
        "Outcome Number", "Outcome", "Program Number", "Program Name",
    ] + [f"{y} EA" for y in YEARS]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).font = openpyxl.styles.Font(bold=True)

    def sort_key(item):
        key, _ = item
        portfolio, agency, outcome_number, _, program_number, program_name = key
        return (
            canon_portfolio(portfolio) or "",
            agency or "",
            outcome_number if outcome_number is not None else -1,
            program_number or "",
            program_name or "",
        )

    for key, g in sorted(groups.items(), key=sort_key):
        (portfolio, agency, outcome_number, outcome_desc,
         program_number, program_name) = key
        canon_agency = g["formal_names"].most_common(1)[0][0] if g["formal_names"] else agency
        row = [
            portfolio, canon_portfolio(portfolio), agency, canon_agency,
            outcome_number, outcome_desc, program_number, program_name,
        ] + [g["years"].get(y) for y in YEARS]
        ws.append(row)

    # Light readability pass: freeze header, widen text columns.
    ws.freeze_panes = "A2"
    widths = [22, 22, 30, 30, 8, 45, 12, 45] + [10] * len(YEARS)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(OUT_PATH)
    print(f"wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
